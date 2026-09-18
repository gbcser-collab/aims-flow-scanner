import io

import fitz
import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont

from aims_docintel.extractor import DocumentExtractor, DocumentPasswordRequired


def _font(size: int = 52):
    for path in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf",
    ):
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            pass
    return ImageFont.load_default()


def test_extracts_native_text_pdf():
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 90), "LOAD Győr 18.09.2026", fontsize=18)
    page.insert_text((72, 125), "UNLOAD Wien 18.09.2026", fontsize=18)
    data = doc.tobytes()
    doc.close()

    result = DocumentExtractor(ocr_languages="eng").extract(data, "order.pdf")
    assert result.kind == "pdf"
    assert "LOAD" in result.full_text
    assert "UNLOAD" in result.full_text
    assert result.pages[0].method in {"embedded-text", "hybrid"}


def test_extracts_docx_text_and_tables():
    doc = Document()
    doc.add_paragraph("Reference: WORD-778899")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "LOAD"
    table.cell(0, 1).text = "Győr"
    table.cell(1, 0).text = "UNLOAD"
    table.cell(1, 1).text = "Brno"
    stream = io.BytesIO()
    doc.save(stream)

    result = DocumentExtractor(ocr_languages="eng").extract(stream.getvalue(), "order.docx")
    assert result.kind == "docx"
    assert "WORD-778899" in result.full_text
    assert "Győr" in result.full_text
    assert "Brno" in result.full_text


def test_scanned_pdf_uses_ocr():
    image = Image.new("RGB", (1800, 1000), "white")
    draw = ImageDraw.Draw(image)
    font = _font()
    draw.text((90, 100), "FREIGHT ORDER", fill="black", font=font)
    draw.text((90, 230), "LOAD BUDAPEST", fill="black", font=font)
    draw.text((90, 360), "UNLOAD VIENNA", fill="black", font=font)
    png = io.BytesIO()
    image.save(png, format="PNG")

    doc = fitz.open()
    page = doc.new_page(width=900, height=500)
    page.insert_image(page.rect, stream=png.getvalue())
    data = doc.tobytes()
    doc.close()

    result = DocumentExtractor(ocr_languages="eng").extract(data, "scan.pdf")
    assert result.pages[0].method == "ocr"
    upper = result.full_text.upper()
    assert "LOAD" in upper
    assert "UNLOAD" in upper


def test_password_protected_pdf_requires_password_and_then_reads():
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 90), "Reference: SECURE-12345", fontsize=18)
    data = doc.tobytes(
        encryption=fitz.PDF_ENCRYPT_AES_256,
        owner_pw="owner-secret",
        user_pw="driver-secret",
        permissions=fitz.PDF_PERM_ACCESSIBILITY,
    )
    doc.close()

    extractor = DocumentExtractor(ocr_languages="eng")
    with pytest.raises(DocumentPasswordRequired):
        extractor.extract(data, "secure.pdf")

    result = extractor.extract(data, "secure.pdf", password="driver-secret")
    assert "SECURE-12345" in result.full_text
