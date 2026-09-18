from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

import fitz
import pikepdf
import pytesseract
from docx import Document
from PIL import Image, ImageFilter, ImageOps
from pypdf import PdfReader


class DocumentIngestError(RuntimeError):
    pass


class DocumentPasswordRequired(DocumentIngestError):
    pass


@dataclass
class ExtractedPage:
    number: int
    text: str
    method: str
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class ExtractedDocument:
    source_name: str
    kind: str
    mime_type: str
    pages: list[ExtractedPage]
    metadata: dict[str, str] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n\n".join(page.text.strip() for page in self.pages if page.text.strip()).strip()

    def to_dict(self, include_page_text: bool = True) -> dict:
        data = {
            "sourceName": self.source_name,
            "kind": self.kind,
            "mimeType": self.mime_type,
            "pageCount": len(self.pages),
            "metadata": self.metadata,
            "warnings": self.warnings,
            "fullText": self.full_text,
            "pages": [page.to_dict() for page in self.pages],
        }
        if not include_page_text:
            for page in data["pages"]:
                page.pop("text", None)
        return data


class DocumentExtractor:
    """
    Layered extractor for freight-order attachments.

    Supported paths:
      - PDF with embedded text
      - scanned/image-only PDF via OCR
      - hybrid PDF, page by page
      - encrypted PDF when a password is provided
      - DOCX/DOCM by reading OOXML without executing macros
      - legacy DOC/RTF/ODT via isolated LibreOffice conversion
      - JPG/PNG/TIFF/WEBP image OCR
      - TXT/CSV fallback

    File extension is never trusted as the primary detector.
    """

    _PDF = b"%PDF"
    _OLE = bytes.fromhex("D0CF11E0A1B11AE1")
    _PNG = bytes.fromhex("89504E470D0A1A0A")
    _JPEG = b"\xff\xd8\xff"
    _TIFF_LE = b"II*\x00"
    _TIFF_BE = b"MM\x00*"
    _RTF = b"{\\rtf"

    def __init__(
        self,
        *,
        ocr_languages: str | None = None,
        max_bytes: int | None = None,
        libreoffice_binary: str | None = None,
    ) -> None:
        self.max_bytes = max_bytes or int(os.getenv("AIMS_DOC_MAX_BYTES", str(25 * 1024 * 1024)))
        self.requested_ocr_languages = (
            ocr_languages
            or os.getenv("AIMS_OCR_LANGS")
            or "hun+eng+deu+pol+slk+ces+fra+lit"
        )
        self.libreoffice_binary = (
            libreoffice_binary
            or os.getenv("AIMS_LIBREOFFICE_BIN")
            or shutil.which("libreoffice")
            or shutil.which("soffice")
        )

    def extract(self, data: bytes, source_name: str, password: str | None = None) -> ExtractedDocument:
        if not data:
            raise DocumentIngestError("A dokumentum üres.")
        if len(data) > self.max_bytes:
            raise DocumentIngestError(
                f"A dokumentum túl nagy ({len(data)} byte). Limit: {self.max_bytes} byte."
            )

        kind = self.detect_kind(data, source_name)
        if kind == "pdf":
            return self._extract_pdf(data, source_name, password=password)
        if kind in {"docx", "docm"}:
            return self._extract_docx(data, source_name, kind=kind)
        if kind in {"doc", "rtf", "odt"}:
            return self._extract_via_libreoffice(data, source_name, kind=kind, password=password)
        if kind in {"png", "jpg", "tiff", "webp"}:
            return self._extract_image(data, source_name, kind=kind)
        if kind in {"txt", "csv"}:
            return self._extract_text(data, source_name, kind=kind)
        raise DocumentIngestError(
            f"Nem támogatott vagy nem felismerhető dokumentumtípus: {source_name}"
        )

    def detect_kind(self, data: bytes, source_name: str = "") -> str:
        head = data[:32]
        lower_name = source_name.lower()
        if head.startswith(self._PDF):
            return "pdf"
        if head.startswith(self._OLE):
            return "doc"
        if head.lstrip().startswith(self._RTF):
            return "rtf"
        if head.startswith(self._PNG):
            return "png"
        if head.startswith(self._JPEG):
            return "jpg"
        if head.startswith(self._TIFF_LE) or head.startswith(self._TIFF_BE):
            return "tiff"
        if head.startswith(b"RIFF") and data[8:12] == b"WEBP":
            return "webp"

        if zipfile.is_zipfile(io.BytesIO(data)):
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as zf:
                    names = set(zf.namelist())
                    if "word/document.xml" in names:
                        return "docm" if lower_name.endswith(".docm") else "docx"
                    if "mimetype" in names:
                        mime = zf.read("mimetype").decode("ascii", "ignore")
                        if "opendocument.text" in mime:
                            return "odt"
            except zipfile.BadZipFile:
                pass

        suffix = Path(lower_name).suffix
        return {
            ".pdf": "pdf",
            ".docx": "docx",
            ".docm": "docm",
            ".doc": "doc",
            ".rtf": "rtf",
            ".odt": "odt",
            ".png": "png",
            ".jpg": "jpg",
            ".jpeg": "jpg",
            ".tif": "tiff",
            ".tiff": "tiff",
            ".webp": "webp",
            ".txt": "txt",
            ".csv": "csv",
        }.get(suffix, "unknown")

    def _extract_pdf(
        self, data: bytes, source_name: str, password: str | None
    ) -> ExtractedDocument:
        warnings: list[str] = []
        pdf_bytes = data

        try:
            doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        except Exception:
            pdf_bytes = self._repair_pdf(data, password=password)
            try:
                doc = fitz.open(stream=pdf_bytes, filetype="pdf")
                warnings.append("A sérült PDF szerkezete automatikusan javítva lett.")
            except Exception as exc:
                raise DocumentIngestError(f"A PDF nem nyitható meg: {exc}") from exc

        try:
            if doc.needs_pass:
                if not password:
                    raise DocumentPasswordRequired("A PDF jelszóval védett.")
                if not doc.authenticate(password):
                    raise DocumentPasswordRequired("A megadott PDF-jelszó hibás.")

            metadata = {
                str(k): str(v)
                for k, v in (doc.metadata or {}).items()
                if v not in (None, "")
            }
            pages: list[ExtractedPage] = []
            fallback_reader = self._safe_pypdf_reader(pdf_bytes, password)

            for index, page in enumerate(doc):
                page_warnings: list[str] = []
                embedded = self._pdf_text(page)
                form_text = self._pdf_form_text(page)
                if form_text:
                    embedded = self._merge_texts(embedded, form_text)

                has_images = bool(page.get_images(full=True))
                text_score = self._text_quality(embedded)
                should_ocr = (
                    text_score < 55
                    or (has_images and len(self._clean_text(embedded)) < 280)
                )

                ocr_text = ""
                if should_ocr:
                    try:
                        ocr_text = self._ocr_pdf_page(page)
                    except Exception as exc:
                        page_warnings.append(f"OCR nem sikerült: {type(exc).__name__}")

                if not embedded.strip() and fallback_reader is not None:
                    try:
                        fallback = fallback_reader.pages[index].extract_text() or ""
                        embedded = self._clean_text(fallback)
                        if embedded:
                            page_warnings.append("PDF szöveg pypdf fallbackból.")
                    except Exception:
                        pass

                combined = self._merge_texts(embedded, ocr_text)
                if not combined.strip():
                    page_warnings.append("Az oldalról nem sikerült olvasható szöveget kinyerni.")

                if embedded.strip() and ocr_text.strip():
                    method = "hybrid"
                elif ocr_text.strip():
                    method = "ocr"
                elif embedded.strip():
                    method = "embedded-text"
                else:
                    method = "empty"

                pages.append(
                    ExtractedPage(
                        number=index + 1,
                        text=combined,
                        method=method,
                        warnings=page_warnings,
                    )
                )

            if not pages:
                raise DocumentIngestError("A PDF nem tartalmaz feldolgozható oldalt.")

            return ExtractedDocument(
                source_name=source_name,
                kind="pdf",
                mime_type="application/pdf",
                pages=pages,
                metadata=metadata,
                warnings=warnings,
            )
        finally:
            doc.close()

    def _repair_pdf(self, data: bytes, password: str | None) -> bytes:
        try:
            with pikepdf.open(io.BytesIO(data), password=password or "") as source:
                out = io.BytesIO()
                source.save(out, linearize=True)
                return out.getvalue()
        except pikepdf.PasswordError as exc:
            raise DocumentPasswordRequired("A PDF jelszóval védett.") from exc
        except Exception as exc:
            raise DocumentIngestError(f"A PDF javítása sikertelen: {exc}") from exc

    def _safe_pypdf_reader(self, data: bytes, password: str | None) -> PdfReader | None:
        try:
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                if not password:
                    return None
                if not reader.decrypt(password):
                    return None
            return reader
        except Exception:
            return None

    def _pdf_text(self, page: fitz.Page) -> str:
        chunks: list[str] = []
        try:
            blocks = page.get_text("blocks", sort=True)
            for block in blocks:
                if len(block) >= 5:
                    text = str(block[4]).strip()
                    if text:
                        chunks.append(text)
        except Exception:
            pass
        if not chunks:
            try:
                chunks.append(page.get_text("text", sort=True))
            except Exception:
                pass
        return self._clean_text("\n".join(chunks))

    def _pdf_form_text(self, page: fitz.Page) -> str:
        rows: list[str] = []
        try:
            widget = page.first_widget
            while widget is not None:
                name = (widget.field_name or "").strip()
                value = str(widget.field_value or "").strip()
                if value:
                    rows.append(f"{name}: {value}" if name else value)
                widget = widget.next
        except Exception:
            return ""
        return "\n".join(rows)

    def _ocr_pdf_page(self, page: fitz.Page) -> str:
        scale = float(os.getenv("AIMS_PDF_OCR_SCALE", "2.4"))
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        image = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
        return self._ocr_image(image)

    def _extract_docx(self, data: bytes, source_name: str, kind: str) -> ExtractedDocument:
        warnings: list[str] = []
        chunks: list[str] = []

        try:
            doc = Document(io.BytesIO(data))
            chunks.extend(p.text for p in doc.paragraphs if p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    values = [cell.text.strip() for cell in row.cells]
                    if any(values):
                        chunks.append(" | ".join(values))
            for section in doc.sections:
                chunks.extend(
                    p.text for p in section.header.paragraphs if p.text.strip()
                )
                chunks.extend(
                    p.text for p in section.footer.paragraphs if p.text.strip()
                )
        except Exception as exc:
            warnings.append(f"python-docx részleges olvasás: {type(exc).__name__}")

        xml_text = self._extract_word_ooxml_text(data)
        combined = self._merge_texts("\n".join(chunks), xml_text)
        if not combined:
            raise DocumentIngestError("A Word dokumentumból nem sikerült szöveget kinyerni.")

        return ExtractedDocument(
            source_name=source_name,
            kind=kind,
            mime_type=(
                "application/vnd.ms-word.document.macroEnabled.12"
                if kind == "docm"
                else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            pages=[ExtractedPage(number=1, text=combined, method="ooxml")],
            warnings=warnings,
        )

    def _extract_word_ooxml_text(self, data: bytes) -> str:
        rows: list[str] = []
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                candidates = [
                    name
                    for name in zf.namelist()
                    if name.startswith("word/")
                    and name.endswith(".xml")
                    and any(
                        token in name
                        for token in (
                            "document.xml",
                            "header",
                            "footer",
                            "footnotes",
                            "endnotes",
                            "comments",
                        )
                    )
                ]
                for name in candidates:
                    try:
                        root = ET.fromstring(zf.read(name))
                    except ET.ParseError:
                        continue
                    for paragraph in root.iter():
                        if not paragraph.tag.endswith("}p"):
                            continue
                        pieces = [
                            node.text or ""
                            for node in paragraph.iter()
                            if node.tag.endswith("}t") and node.text
                        ]
                        text = "".join(pieces).strip()
                        if text:
                            rows.append(text)
        except Exception:
            return ""
        return self._clean_text("\n".join(rows))

    def _extract_via_libreoffice(
        self,
        data: bytes,
        source_name: str,
        kind: str,
        password: str | None,
    ) -> ExtractedDocument:
        if not self.libreoffice_binary:
            raise DocumentIngestError(
                "A régi Word/RTF/ODT feldolgozásához LibreOffice szükséges a szerveren."
            )

        suffix = Path(source_name).suffix or f".{kind}"
        with tempfile.TemporaryDirectory(prefix="aims-docintel-") as temp:
            temp_path = Path(temp)
            input_path = temp_path / f"input{suffix}"
            output_dir = temp_path / "out"
            profile_dir = temp_path / "lo-profile"
            output_dir.mkdir()
            profile_dir.mkdir()
            input_path.write_bytes(data)

            profile_uri = profile_dir.resolve().as_uri()
            command = [
                self.libreoffice_binary,
                "--headless",
                "--norestore",
                "--nodefault",
                "--nolockcheck",
                f"-env:UserInstallation={profile_uri}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_dir),
                str(input_path),
            ]
            try:
                result = subprocess.run(
                    command,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    timeout=45,
                    check=False,
                )
            except subprocess.TimeoutExpired as exc:
                raise DocumentIngestError("A Word konverzió időtúllépés miatt megszakadt.") from exc

            candidates = list(output_dir.glob("*.pdf"))
            if result.returncode != 0 or not candidates:
                stderr = result.stderr.decode("utf-8", "ignore")[-500:]
                raise DocumentIngestError(f"LibreOffice konverziós hiba: {stderr}")

            extracted = self._extract_pdf(
                candidates[0].read_bytes(),
                source_name=source_name,
                password=password,
            )
            extracted.kind = kind
            extracted.mime_type = {
                "doc": "application/msword",
                "rtf": "application/rtf",
                "odt": "application/vnd.oasis.opendocument.text",
            }[kind]
            extracted.warnings.insert(
                0, f"{kind.upper()} dokumentum PDF-re normalizálva, majd feldolgozva."
            )
            return extracted

    def _extract_image(self, data: bytes, source_name: str, kind: str) -> ExtractedDocument:
        try:
            image = Image.open(io.BytesIO(data)).convert("RGB")
        except Exception as exc:
            raise DocumentIngestError(f"A kép nem olvasható: {exc}") from exc
        text = self._ocr_image(image)
        return ExtractedDocument(
            source_name=source_name,
            kind=kind,
            mime_type={
                "png": "image/png",
                "jpg": "image/jpeg",
                "tiff": "image/tiff",
                "webp": "image/webp",
            }[kind],
            pages=[ExtractedPage(number=1, text=text, method="ocr")],
        )

    def _extract_text(self, data: bytes, source_name: str, kind: str) -> ExtractedDocument:
        text = self._decode_text(data)
        if not text.strip():
            raise DocumentIngestError("A szöveges dokumentum üres.")
        return ExtractedDocument(
            source_name=source_name,
            kind=kind,
            mime_type="text/csv" if kind == "csv" else "text/plain",
            pages=[ExtractedPage(number=1, text=self._clean_text(text), method="text")],
        )

    def _ocr_image(self, image: Image.Image) -> str:
        langs = self._available_ocr_languages()
        gray = ImageOps.grayscale(image)
        gray = ImageOps.autocontrast(gray)
        gray = gray.filter(ImageFilter.MedianFilter(size=3))

        candidates: list[str] = []
        for psm in (6, 11):
            try:
                candidates.append(
                    pytesseract.image_to_string(
                        gray,
                        lang=langs,
                        config=f"--oem 3 --psm {psm}",
                    )
                )
            except Exception:
                pass

        best = max(candidates, key=self._text_quality, default="")
        if self._text_quality(best) < 60:
            for angle in (90, 180, 270):
                rotated = gray.rotate(angle, expand=True)
                try:
                    value = pytesseract.image_to_string(
                        rotated,
                        lang=langs,
                        config="--oem 3 --psm 6",
                    )
                except Exception:
                    continue
                if self._text_quality(value) > self._text_quality(best):
                    best = value

        return self._clean_text(best)

    def _available_ocr_languages(self) -> str:
        requested = [item for item in self.requested_ocr_languages.split("+") if item]
        try:
            installed = set(pytesseract.get_languages(config=""))
        except Exception as exc:
            raise DocumentIngestError(
                "Tesseract OCR nem érhető el a dokumentumfeldolgozó szerveren."
            ) from exc

        selected = [language for language in requested if language in installed]
        if not selected and "eng" in installed:
            selected = ["eng"]
        if not selected:
            raise DocumentIngestError("Nincs használható Tesseract nyelvi csomag telepítve.")
        return "+".join(selected)

    def _merge_texts(self, first: str, second: str) -> str:
        first = self._clean_text(first)
        second = self._clean_text(second)
        if not first:
            return second
        if not second:
            return first

        first_lines = {self._line_key(line) for line in first.splitlines() if line.strip()}
        extra: list[str] = []
        for line in second.splitlines():
            key = self._line_key(line)
            if key and key not in first_lines:
                extra.append(line)
                first_lines.add(key)
        return self._clean_text(first + ("\n" + "\n".join(extra) if extra else ""))

    def _line_key(self, value: str) -> str:
        return re.sub(r"[^a-z0-9]+", "", value.lower())[:180]

    def _clean_text(self, value: str) -> str:
        value = value.replace("\x00", " ").replace("\r", "\n")
        value = re.sub(r"[\t\f\v]+", " ", value)
        value = re.sub(r" +", " ", value)
        value = re.sub(r"\n{3,}", "\n\n", value)
        return "\n".join(line.strip() for line in value.splitlines()).strip()

    def _decode_text(self, data: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-16", "cp1250", "latin-1"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", "replace")

    def _text_quality(self, value: str) -> int:
        if not value:
            return 0
        compact = re.sub(r"\s+", "", value)
        if not compact:
            return 0
        alnum = sum(ch.isalnum() for ch in compact)
        letters = sum(ch.isalpha() for ch in compact)
        separators = sum(ch in ":,.-/()" for ch in value)
        return min(1000, alnum + letters + min(separators, 80))
